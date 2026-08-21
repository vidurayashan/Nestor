import pdfplumber
from docx import Document
from PIL import Image
import io
from pathlib import Path
from typing import List, Tuple, Dict, Any
import ffmpeg
import json
import subprocess
import tempfile


class ProcessedReport:
    """Container for processed report data"""
    def __init__(self):
        self.text: str = ""
        self.images: List[Image.Image] = []
        self.page_count: int = 0


class ProcessedVideo:
    """Container for processed video data"""
    def __init__(self):
        self.transcript: str = ""
        self.frames: List[Image.Image] = []
        self.duration: float = 0  # seconds
        self.audio_path: Optional[Path] = None


def extract_pdf_content(pdf_path: Path) -> ProcessedReport:
    """Extract text and images from PDF"""
    result = ProcessedReport()
    
    with pdfplumber.open(pdf_path) as pdf:
        result.page_count = len(pdf.pages)
        
        # Extract text
        text_parts = []
        for page in pdf.pages:
            text = page.extract_text()
            if text:
                text_parts.append(text)
        
        result.text = "\n\n".join(text_parts)
        
        # Extract images
        for page_num, page in enumerate(pdf.pages):
            # Get images from page
            if hasattr(page, 'images'):
                for img_info in page.images:
                    try:
                        # Extract image using pdfplumber's method
                        img_obj = page.within_bbox((
                            img_info['x0'], img_info['top'],
                            img_info['x1'], img_info['bottom']
                        )).to_image()
                        
                        # Convert to PIL Image
                        pil_img = img_obj.original
                        result.images.append(pil_img)
                    except Exception as e:
                        print(f"Warning: Could not extract image from page {page_num + 1}: {e}")
    
    return result


def extract_docx_content(docx_path: Path) -> ProcessedReport:
    """Extract text and images from Word document"""
    result = ProcessedReport()
    
    doc = Document(docx_path)
    
    # Approximate page count (assuming ~500 words per page)
    word_count = sum(len(para.text.split()) for para in doc.paragraphs)
    result.page_count = max(1, word_count // 500)
    
    # Extract text
    text_parts = [para.text for para in doc.paragraphs if para.text.strip()]
    result.text = "\n\n".join(text_parts)
    
    # Extract images
    for rel in doc.part.rels.values():
        if "image" in rel.target_ref:
            try:
                image_data = rel.target_part.blob
                image = Image.open(io.BytesIO(image_data))
                result.images.append(image)
            except Exception as e:
                print(f"Warning: Could not extract embedded image: {e}")
    
    return result


def process_report(file_path: Path) -> ProcessedReport:
    """Process report file (PDF or docx)"""
    suffix = file_path.suffix.lower()
    
    if suffix == '.pdf':
        return extract_pdf_content(file_path)
    elif suffix in ['.docx', '.doc']:
        return extract_docx_content(file_path)
    else:
        raise ValueError(f"Unsupported report format: {suffix}")


def get_video_duration(video_path: Path) -> float:
    """Get video duration in seconds using ffprobe"""
    try:
        probe = ffmpeg.probe(str(video_path))
        duration = float(probe['streams'][0]['duration'])
        return duration
    except Exception as e:
        print(f"Warning: Could not get video duration: {e}")
        # Fallback: try with ffprobe directly
        try:
            result = subprocess.run(
                ['ffprobe', '-v', 'error', '-show_entries', 
                 'format=duration', '-of', 
                 'default=noprint_wrappers=1:nokey=1', str(video_path)],
                capture_output=True,
                text=True
            )
            return float(result.stdout.strip())
        except:
            return 0


def extract_audio_from_video(video_path: Path, output_path: Path) -> Path:
    """Extract audio track from video for transcription"""
    try:
        (
            ffmpeg
            .input(str(video_path))
            .output(str(output_path), acodec='pcm_s16le', ac=1, ar='16k')
            .overwrite_output()
            .run(capture_output=True)
        )
        return output_path
    except ffmpeg.Error as e:
        raise RuntimeError(f"Failed to extract audio: {e.stderr.decode()}")


def sample_video_frames_interval(
    video_path: Path, 
    interval: int = 25,
    max_frames: int = 50
) -> List[Image.Image]:
    """
    Sample frames from video at regular intervals.
    
    Args:
        video_path: Path to video file
        interval: Seconds between frames
        max_frames: Maximum number of frames to extract
    """
    frames = []
    duration = get_video_duration(video_path)
    
    if duration == 0:
        return frames
    
    # Calculate actual interval to not exceed max_frames
    num_possible_frames = int(duration / interval)
    if num_possible_frames > max_frames:
        interval = duration / max_frames
    
    timestamps = []
    current = 0
    while current < duration and len(timestamps) < max_frames:
        timestamps.append(current)
        current += interval
    
    # Extract frames at timestamps
    for i, timestamp in enumerate(timestamps):
        try:
            out, _ = (
                ffmpeg
                .input(str(video_path), ss=timestamp)
                .filter('select', f'gte(n,{i})')
                .output('pipe:', vframes=1, format='image2', vcodec='png')
                .run(capture_output=True)
            )
            
            image = Image.open(io.BytesIO(out))
            frames.append(image)
        except Exception as e:
            print(f"Warning: Could not extract frame at {timestamp}s: {e}")
    
    return frames


def sample_video_frames_scene_change(
    video_path: Path,
    max_frames: int = 50,
    threshold: float = 0.3
) -> List[Image.Image]:
    """
    Sample frames from video using scene change detection.
    
    Args:
        video_path: Path to video file
        max_frames: Maximum number of frames to extract
        threshold: Scene change detection threshold (0.0-1.0)
    """
    frames = []
    
    try:
        # Use ffmpeg scene detection filter
        with tempfile.TemporaryDirectory() as tmpdir:
            tmpdir_path = Path(tmpdir)
            
            # Run ffmpeg with scene detection
            (
                ffmpeg
                .input(str(video_path))
                .filter('select', f'gt(scene,{threshold})')
                .output(
                    str(tmpdir_path / 'frame_%04d.png'),
                    vsync='vfr',
                    vframes=max_frames
                )
                .overwrite_output()
                .run(capture_output=True)
            )
            
            # Load extracted frames
            for frame_file in sorted(tmpdir_path.glob('frame_*.png')):
                image = Image.open(frame_file)
                frames.append(image.copy())
                
                if len(frames) >= max_frames:
                    break
    
    except Exception as e:
        print(f"Warning: Scene change detection failed, falling back to interval sampling: {e}")
        # Fallback to interval-based sampling
        return sample_video_frames_interval(video_path, interval=25, max_frames=max_frames)
    
    return frames
